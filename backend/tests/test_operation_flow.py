"""End-to-end operation flow tests."""
import sys
import os
import time
from io import BytesIO

import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import StaticPool

sys.path.insert(0, os.path.join(os.path.dirname(__file__), '..'))

from app.database import Base, get_db
from app.main import app
from app.services.lookup_cache import lookup_cache


def stock_excel() -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.append(["ETİKET", "REFERANS", "MİKTAR", "X98FIFO_TARIH"])
    rows = [
        ["A001", "REF-X", 30, "09.07.2026 18:00"],
        ["A002", "REF-X", 30, "10.07.2026 08:15"],
        ["A003", "REF-X", 30, "10.07.2026 09:30"],
        ["A004", "REF-X", 30, "10.07.2026 14:20"],
        ["A005", "REF-X", 30, "11.07.2026 08:00"],
        ["B001", "OTHER", 50, "01.08.2026 10:00"],
    ]
    for r in rows:
        ws.append(r)
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


@pytest.fixture
def client():
    engine = create_engine("sqlite://", connect_args={"check_same_thread": False}, poolclass=StaticPool)
    Base.metadata.create_all(engine)
    TestSession = sessionmaker(bind=engine)

    lookup_cache._pools.clear()
    lookup_cache._label_data.clear()
    lookup_cache._ref_context.clear()
    lookup_cache._all_labels.clear()
    lookup_cache._global_label_shipment.clear()

    import app.database as db_module
    db_module.SessionLocal = TestSession

    def override_get_db():
        db = TestSession()
        try:
            yield db
        finally:
            db.close()

    app.dependency_overrides[get_db] = override_get_db
    with TestClient(app) as c:
        yield c
    app.dependency_overrides.clear()


class TestOperationFlow:
    def test_full_operation_flow(self, client):
        r = client.post("/api/inventory/import/stock", files={
            "file": ("s.xlsx", stock_excel(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        })
        assert r.status_code == 200

        r = client.post("/api/shipment/targets", json={"reference": "REF-X", "target_quantity": 90})
        assert r.status_code == 200

        r = client.post("/api/shipment/find")
        assert r.status_code == 200
        assert len(r.json()["shipments"]) == 1

        r = client.post("/api/shipment/scan", json={"label": "A002"})
        assert r.json()["result"] == "SEVKİYAT ÜRÜNÜ"

        r = client.post("/api/shipment/scan", json={"label": "A002"})
        assert r.json()["result"] == "ZATEN OKUTULDU"

        r = client.post("/api/shipment/scan", json={"label": "A005"})
        assert r.json()["result"] == "SEVKİYAT DIŞI"

        r = client.post("/api/shipment/scan", json={"label": "ZZZZZ"})
        assert r.json()["result"] == "ETİKET BULUNAMADI"

        r = client.post("/api/shipment/scan", json={"label": "B001"})
        assert r.json()["result"] == "SEVKİYAT DIŞI"

    def test_manual_equals_excel_target(self, client):
        client.post("/api/inventory/import/stock", files={
            "file": ("s.xlsx", stock_excel(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        })

        wb = Workbook()
        ws = wb.active
        ws.append(["REFERANS", "MİKTAR"])
        ws.append(["REF-X", 90])
        buf = BytesIO()
        wb.save(buf)

        r = client.post("/api/shipment/targets/import", files={
            "file": ("t.xlsx", buf.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        })
        assert r.status_code == 200

        r = client.post("/api/shipment/find")
        # Hedef 90, tam FIFO sırasıyla havuz 90 (09.07 + 10.07'den gerektiği kadar)
        assert r.json()["shipments"][0]["pool_quantity"] == 90
        assert r.json()["shipments"][0]["requested_quantity"] == 90

    def test_reset_and_undo_scan(self, client):
        client.post("/api/inventory/import/stock", files={
            "file": ("s.xlsx", stock_excel(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        })
        client.post("/api/shipment/targets", json={"reference": "REF-X", "target_quantity": 90})
        find = client.post("/api/shipment/find")
        shipment_id = find.json()["shipments"][0]["shipment_id"]

        client.post("/api/shipment/scan", json={"label": "A001"})
        client.post("/api/shipment/scan", json={"label": "A002"})

        scanned = client.get(f"/api/shipment/{shipment_id}/scanned")
        assert scanned.status_code == 200
        assert len(scanned.json()) == 2

        undo = client.delete(f"/api/shipment/{shipment_id}/scans/A001")
        assert undo.status_code == 200
        assert undo.json()["scanned_quantity"] == 30

        scanned2 = client.get(f"/api/shipment/{shipment_id}/scanned")
        assert len(scanned2.json()) == 1

        rescan = client.post("/api/shipment/scan", json={"label": "A001"})
        assert rescan.json()["result"] == "SEVKİYAT ÜRÜNÜ"

        reset = client.post("/api/shipment/reset")
        assert reset.status_code == 200
        assert reset.json()["cancelled"] == 1

        status = client.get("/api/shipment/status")
        assert status.json() == []

    def test_single_target_can_be_removed(self, client):
        first = client.post("/api/shipment/targets", json={"reference": "REF-A", "target_quantity": 15})
        second = client.post("/api/shipment/targets", json={"reference": "REF-B", "target_quantity": 30})
        assert first.status_code == 200
        assert second.status_code == 200

        removed = client.delete(f"/api/shipment/targets/{first.json()['id']}")
        assert removed.status_code == 200

        remaining = client.get("/api/shipment/targets")
        assert remaining.status_code == 200
        assert [row["reference"] for row in remaining.json()] == ["REF-B"]

        missing = client.delete("/api/shipment/targets/999999")
        assert missing.status_code == 404

    def test_recalculate_uses_current_stock_and_keeps_scans(self, client):
        def stock_file(rows):
            wb = Workbook()
            ws = wb.active
            ws.append(["ETİKET", "REFERANS", "MİKTAR", "X98FIFO_TARIH"])
            for row in rows:
                ws.append(row)
            buf = BytesIO()
            wb.save(buf)
            return buf.getvalue()

        first = stock_file([
            ["A001", "REF-X", 30, "01.06.2026 08:15"],
            ["A002", "REF-X", 30, "02.06.2026 09:30"],
            ["A003", "REF-X", 30, "03.06.2026 10:45"],
        ])
        uploaded = client.post("/api/inventory/import/stock", files={
            "file": ("s.xlsx", first, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        })
        assert uploaded.status_code == 200

        client.post("/api/shipment/targets", json={"reference": "REF-X", "target_quantity": 60})
        found = client.post("/api/shipment/find")
        assert found.status_code == 200
        shipment_id = found.json()["shipments"][0]["shipment_id"]

        scan = client.post("/api/shipment/scan", json={"label": "A001"})
        assert scan.status_code == 200
        assert scan.json()["result"] == "SEVKİYAT ÜRÜNÜ"

        extra = stock_file([
            ["A000", "REF-X", 30, "20.05.2026 08:15"],
        ])
        appended = client.post(
            "/api/inventory/import/stock?replace=false",
            files={"file": ("s2.xlsx", extra, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        )
        assert appended.status_code == 200

        recalc = client.post("/api/shipment/recalculate")
        assert recalc.status_code == 200
        assert recalc.json()["status"] == "started"
        status = {"phase": "running"}
        for _ in range(100):
            status = client.get("/api/shipment/recalculate/status").json()
            if status["phase"] in ("done", "error"):
                break
            time.sleep(0.05)
        assert status["phase"] == "done", status
        assert status["updated"] == 1

        manifest = client.get("/api/shipment/manifest").json()
        item = next(row for row in manifest if row["shipment_id"] == shipment_id)
        dates = {row["fifo_group_date"]: row for row in item["items"]}
        assert item["scanned_quantity"] == 30
        assert "20.05.2026" in dates
        assert dates["20.05.2026"]["quantity"] == 30
        assert dates["01.06.2026"]["scanned_quantity"] == 30
        assert dates["01.06.2026"]["scanned_labels"][0]["label"] == "A001"
        assert dates["20.05.2026"]["outside_fifo"] is False
        assert dates["01.06.2026"]["outside_fifo"] is False
        assert dates["01.06.2026"]["scanned_labels"][0]["outside_fifo"] is False
        assert "02.06.2026" not in dates

    def test_scan_outside_recalculated_window_is_marked(self, client):
        def stock_file(rows):
            wb = Workbook()
            ws = wb.active
            ws.append(["ETİKET", "REFERANS", "MİKTAR", "X98FIFO_TARIH"])
            for row in rows:
                ws.append(row)
            buf = BytesIO()
            wb.save(buf)
            return buf.getvalue()

        def wait_recalc():
            started = client.post("/api/shipment/recalculate")
            assert started.status_code == 200
            status = {"phase": "running"}
            for _ in range(100):
                status = client.get("/api/shipment/recalculate/status").json()
                if status["phase"] in ("done", "error"):
                    break
                time.sleep(0.05)
            assert status["phase"] == "done", status

        uploaded = client.post("/api/inventory/import/stock", files={
            "file": ("s.xlsx", stock_file([
                ["L1", "REF-X", 30, "01.06.2026 08:15"],
                ["L2", "REF-X", 30, "02.06.2026 09:30"],
                ["L7", "REF-X", 30, "07.06.2026 10:45"],
            ]), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        })
        assert uploaded.status_code == 200

        client.post("/api/shipment/targets", json={"reference": "REF-X", "target_quantity": 30})
        assert client.post("/api/shipment/find").status_code == 200
        client.post("/api/shipment/targets", json={"reference": "REF-X", "target_quantity": 60})
        assert client.post("/api/shipment/find").status_code == 200

        groups = client.get("/api/shipment/groups").json()
        early = next(g for g in groups if g["requested_quantity"] == 30)
        later = next(g for g in groups if g["requested_quantity"] == 60)

        scan = client.post("/api/shipment/scan", json={"label": "L7", "group_id": later["group_id"]})
        assert scan.status_code == 200
        assert scan.json()["result"] == "SEVKİYAT ÜRÜNÜ"

        removed = client.delete(f"/api/shipment/group/{early['group_id']}")
        assert removed.status_code == 200
        wait_recalc()

        manifest = client.get("/api/shipment/manifest").json()
        item = next(row for row in manifest if row["shipment_id"] == later["items"][0]["shipment_id"])
        dates = {row["fifo_group_date"]: row for row in item["items"]}
        assert dates["01.06.2026"]["outside_fifo"] is False
        assert dates["07.06.2026"]["outside_fifo"] is True
        assert dates["07.06.2026"]["outside_reason"] == "window"
        assert dates["07.06.2026"]["scanned_labels"][0]["label"] == "L7"
        assert dates["07.06.2026"]["scanned_labels"][0]["outside_fifo"] is True

        inside = client.post("/api/shipment/scan", json={"label": "L1", "group_id": later["group_id"]})
        assert inside.status_code == 200
        assert inside.json()["result"] == "SEVKİYAT ÜRÜNÜ"
        manifest = client.get("/api/shipment/manifest").json()
        item = next(row for row in manifest if row["shipment_id"] == later["items"][0]["shipment_id"])
        dates = {row["fifo_group_date"]: row for row in item["items"]}
        assert dates["01.06.2026"]["scanned_labels"][0]["outside_fifo"] is False
        assert dates["07.06.2026"]["scanned_labels"][0]["outside_fifo"] is True

    def test_fifo_date_lists_labels_that_can_still_be_scanned(self, client):
        wb = Workbook()
        ws = wb.active
        ws.append(["ETİKET", "REFERANS", "MİKTAR", "X98FIFO_TARIH"])
        ws.append(["L1", "REF-X", 15, "01.06.2026 08:15"])
        ws.append(["L2", "REF-X", 15, "01.06.2026 09:30"])
        ws.append(["L3", "REF-X", 15, "02.06.2026 10:45"])
        buf = BytesIO()
        wb.save(buf)
        uploaded = client.post("/api/inventory/import/stock", files={
            "file": ("s.xlsx", buf.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        })
        assert uploaded.status_code == 200
        client.post("/api/shipment/targets", json={"reference": "REF-X", "target_quantity": 30})
        found = client.post("/api/shipment/find")
        assert found.status_code == 200
        shipment_id = found.json()["shipments"][0]["shipment_id"]

        open_day = client.get(f"/api/shipment/{shipment_id}/candidates", params={"fifo_group": "01.06.2026"})
        assert open_day.status_code == 200
        body = open_day.json()
        assert body["remaining"] == 30
        assert [row["label"] for row in body["labels"]] == ["L1", "L2"]

        later = client.get(f"/api/shipment/{shipment_id}/candidates", params={"fifo_group": "02.06.2026"})
        assert later.status_code == 200
        assert later.json()["labels"] == []

        scan = client.post("/api/shipment/scan", json={"label": "L1"})
        assert scan.json()["result"] == "SEVKİYAT ÜRÜNÜ"
        after = client.get(f"/api/shipment/{shipment_id}/candidates", params={"fifo_group": "01.06.2026"}).json()
        assert [row["label"] for row in after["labels"]] == ["L2"]
        assert after["remaining"] == 15

    def test_headerless_shipment_excel_names_the_shipment(self, client):
        client.post("/api/inventory/import/stock", files={
            "file": ("s.xlsx", stock_excel(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        })

        wb = Workbook()
        ws = wb.active
        ws["A1"] = "Genel rapor"
        ws["C3"] = "Not"
        ws["E6"] = "REF-X"
        ws["F6"] = 30
        ws["E7"] = "Başlık satırı"
        buf = BytesIO()
        wb.save(buf)

        imported = client.post("/api/shipment/targets/import", files={
            "file": (
                "Merkez Depo.xlsx",
                buf.getvalue(),
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        })
        assert imported.status_code == 200, imported.text
        body = imported.json()
        assert body["successful"] == 1
        assert body["targets"][0]["reference"] == "REF-X"
        assert body["targets"][0]["target_quantity"] == 30

        found = client.post("/api/shipment/find")
        assert found.status_code == 200
        groups = client.get("/api/shipment/groups").json()
        assert groups[0]["name"] == "Merkez Depo"

    def test_text_instead_of_quantity_is_not_a_shipment_file(self, client):
        kept = client.post("/api/shipment/targets", json={"reference": "REF-KEEP", "target_quantity": 12})
        assert kept.status_code == 200

        wb = Workbook()
        ws = wb.active
        ws["B4"] = "6036003-G"
        ws["C4"] = "Depo"
        ws["B5"] = "6036004-G"
        ws["C5"] = "Açık"
        buf = BytesIO()
        wb.save(buf)

        imported = client.post("/api/shipment/targets/import", files={
            "file": (
                "yanlis.xlsx",
                buf.getvalue(),
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        })
        assert imported.status_code == 400
        assert "sevkiyat formatına uygun değil" in imported.json()["detail"]

        remaining = client.get("/api/shipment/targets").json()
        assert [row["reference"] for row in remaining] == ["REF-KEEP"]
        assert client.get("/api/shipment/groups").json() == []
