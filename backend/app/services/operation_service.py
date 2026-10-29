from datetime import datetime, timedelta
from decimal import Decimal

from sqlalchemy.orm import Session, joinedload

from app.models import (
    Shipment, ShipmentLabel, ShipmentStatus, ShipmentLabelStatus,
    ShipmentDateQuota, ScanLog, ScanResult, InventoryLabel,
)
from app.services.fifo_quota import format_group, group_key, quota_map, scanned_on_date
from app.services.lookup_cache import lookup_cache
from app.services.shipment_service import (
    _build_shipments_progress_batch,
    get_shipment_progress,
    ideal_fifo_dates,
    rebalance_reference_pending,
)


def delete_shipment_group(db: Session, group_id: int) -> int:
    """Tek bir sevkiyat grubunu, etiket ve okutma kayıtlarıyla birlikte siler.

    Diğer sevkiyat gruplarına ve hedef listesine dokunmaz.
    """
    shipments = db.query(Shipment).filter(Shipment.group_id == group_id).all()
    if not shipments:
        # Eski kayıtlarda grup numarası yoktu; grup kimliği sevkiyatın kendi id'siydi.
        # Gerçek bir grup varken aynı numaralı eski kaydı silme.
        shipments = db.query(Shipment).filter(
            Shipment.group_id.is_(None),
            Shipment.id == group_id,
        ).all()
    if not shipments:
        raise ValueError("Sevkiyat grubu bulunamadı")

    shipment_ids = [s.id for s in shipments]
    db.query(ShipmentDateQuota).filter(
        ShipmentDateQuota.shipment_id.in_(shipment_ids)
    ).delete(synchronize_session=False)
    db.query(ShipmentLabel).filter(
        ShipmentLabel.shipment_id.in_(shipment_ids)
    ).delete(synchronize_session=False)
    db.query(ScanLog).filter(
        ScanLog.shipment_id.in_(shipment_ids)
    ).delete(synchronize_session=False)
    db.query(Shipment).filter(
        Shipment.id.in_(shipment_ids)
    ).delete(synchronize_session=False)
    db.commit()

    for sid in shipment_ids:
        lookup_cache.unload_shipment(sid)
    lookup_cache.rebuild_global_index()
    return len(shipment_ids)


def reset_active_shipments(db: Session) -> int:
    """Aktif ve tamamlanmış tüm sevkiyatları SİL ve DB'yi temizle (VPS depolama şişmesin)."""
    from app.models import ShipmentTarget

    active = db.query(Shipment).filter(
        Shipment.status.in_([ShipmentStatus.ACTIVE, ShipmentStatus.COMPLETED])
    ).all()
    count = len(active)
    shipment_ids = [s.id for s in active]

    if shipment_ids:
        # 1. Sevkiyatlara bağlı etiketleri sil
        db.query(ShipmentDateQuota).filter(
            ShipmentDateQuota.shipment_id.in_(shipment_ids)
        ).delete(synchronize_session=False)

        db.query(ShipmentLabel).filter(
            ShipmentLabel.shipment_id.in_(shipment_ids)
        ).delete(synchronize_session=False)

        # 2. Sevkiyatlara bağlı scan loglarını sil
        db.query(ScanLog).filter(
            ScanLog.shipment_id.in_(shipment_ids)
        ).delete(synchronize_session=False)

        # 3. Sevkiyat kayıtlarını sil
        db.query(Shipment).filter(
            Shipment.id.in_(shipment_ids)
        ).delete(synchronize_session=False)

    # 4. Sevkiyata bağlı olmayan orphan scan loglarını da sil
    db.query(ScanLog).filter(
        ScanLog.shipment_id.is_(None)
    ).delete(synchronize_session=False)

    # 5. Hedef listesini temizle
    db.query(ShipmentTarget).delete(synchronize_session=False)

    db.commit()

    # 6. Cache'i tamamen temizle
    for sid in shipment_ids:
        lookup_cache.unload_shipment(sid)
    lookup_cache.rebuild_global_index()

    # 7. SQLite VACUUM — disk alanını geri kazan
    try:
        raw_conn = db.bind.raw_connection()
        raw_conn.execute("VACUUM")
        raw_conn.close()
    except Exception:
        pass  # PostgreSQL veya farklı DB'lerde hata yutulur

    return count


def get_scanned_labels(db: Session, shipment_id: int) -> list[dict]:
    shipment = db.query(Shipment).filter(Shipment.id == shipment_id).first()
    if not shipment:
        raise ValueError("Sevkiyat bulunamadı")

    sls = (
        db.query(ShipmentLabel)
        .options(joinedload(ShipmentLabel.inventory_label))
        .filter(
            ShipmentLabel.shipment_id == shipment_id,
            ShipmentLabel.status == ShipmentLabelStatus.SCANNED,
        )
        .all()
    )

    result = []
    for sl in sls:
        inv = sl.inventory_label
        if not inv:
            continue
        log = (
            db.query(ScanLog)
            .filter(
                ScanLog.shipment_id == shipment_id,
                ScanLog.scanned_value == inv.label,
                ScanLog.result == ScanResult.SHIPMENT_PRODUCT,
            )
            .order_by(ScanLog.scanned_at.desc())
            .first()
        )
        result.append({
            "label": inv.label,
            "quantity": float(sl.allocated_quantity),
            "fifo_date": inv.fifo_date.strftime("%d.%m.%Y"),
            "scanned_at": log.scanned_at.isoformat() if log else None,
        })

    result.sort(key=lambda x: x.get("scanned_at") or "", reverse=True)
    return result


from app.services.pool_validation import find_inventory_label


def undo_scan(db: Session, shipment_id: int, label: str) -> dict:
    shipment = db.query(Shipment).filter(Shipment.id == shipment_id).first()
    if not shipment:
        raise ValueError("Sevkiyat bulunamadı")

    if shipment.status == ShipmentStatus.CANCELLED:
        raise ValueError("Sevkiyat iptal edilmiş")

    inv = find_inventory_label(db, label)
    target_label = inv.label if inv else label.strip()

    sl = (
        db.query(ShipmentLabel)
        .join(InventoryLabel)
        .filter(
            ShipmentLabel.shipment_id == shipment_id,
            InventoryLabel.label == target_label,
        )
        .first()
    )

    if not sl or sl.status != ShipmentLabelStatus.SCANNED:
        raise ValueError("Bu etiket okutulmuş olarak kayıtlı değil")

    was_completed = shipment.status == ShipmentStatus.COMPLETED

    sl.scanned_quantity = 0
    sl.status = ShipmentLabelStatus.PENDING
    db.flush()

    if was_completed:
        shipment.status = ShipmentStatus.ACTIVE
        shipment.completed_at = None

    db.query(ScanLog).filter(
        ScanLog.shipment_id == shipment_id,
        ScanLog.scanned_value == label.strip(),
        ScanLog.result == ScanResult.SHIPMENT_PRODUCT,
    ).delete()

    rebalance_reference_pending(db, shipment.reference)
    db.commit()

    if shipment.status == ShipmentStatus.ACTIVE:
        if was_completed:
            lookup_cache.load_shipment(db, shipment_id)
        else:
            lookup_cache.unmark_scanned(shipment_id, label.strip())

    return get_shipment_progress(db, shipment_id)


def get_shipment_manifest(
    db: Session,
    group_id: int | None = None,
    shipment_id: int | None = None,
) -> list[dict]:
    """
    Aktif ve tamamlanmış sevkiyatlar için sistemin arka planda hesapladığı
    tam FIFO havuzunu, etiket listesini ve okutma durumlarını toplu (batch) olarak döndürür.
    """
    from collections import defaultdict
    from sqlalchemy import func

    query = db.query(Shipment).filter(
        Shipment.status.in_([ShipmentStatus.ACTIVE, ShipmentStatus.COMPLETED])
    )
    if shipment_id is not None:
        query = query.filter(Shipment.id == shipment_id)
    elif group_id is not None:
        query = query.filter((Shipment.group_id == group_id) | (Shipment.id == group_id))

    shipments = query.order_by(
        Shipment.sort_order.asc().nulls_last(),
        Shipment.created_at.asc(),
        Shipment.id.asc(),
    ).all()
    if not shipments:
        return []

    sids = [s.id for s in shipments]
    progress_list = _build_shipments_progress_batch(db, shipments)
    prog_by_sid = {p["shipment_id"]: p for p in progress_list}

    # 1. Tüm tarih kotalarını tek sorguda çek
    quota_rows = (
        db.query(ShipmentDateQuota)
        .filter(ShipmentDateQuota.shipment_id.in_(sids))
        .all()
    )
    quotas_by_sid: dict[int, dict[datetime, Decimal]] = defaultdict(dict)
    hourly_by_sid = {s.id: bool(getattr(s, "hourly_fifo", False)) for s in shipments}
    for qrow in quota_rows:
        h = hourly_by_sid.get(qrow.shipment_id, False)
        k = group_key(qrow.fifo_group, h)
        quotas_by_sid[qrow.shipment_id][k] = Decimal(qrow.quota_quantity)

    # 2. Tüm okutulan etiketleri tek sorguda çek
    scanned_rows = (
        db.query(ShipmentLabel)
        .options(joinedload(ShipmentLabel.inventory_label))
        .filter(
            ShipmentLabel.shipment_id.in_(sids),
            ShipmentLabel.status.in_([ShipmentLabelStatus.SCANNED, ShipmentLabelStatus.PARTIAL]),
        )
        .all()
    )
    scanned_by_sid_date: dict[int, dict[datetime, list[ShipmentLabel]]] = defaultdict(lambda: defaultdict(list))
    scanned_sum_by_sid_date: dict[int, dict[datetime, Decimal]] = defaultdict(lambda: defaultdict(lambda: Decimal("0")))
    refs_with_scans: set[str] = set()
    for sl in scanned_rows:
        inv = sl.inventory_label
        if not inv:
            continue
        h = hourly_by_sid.get(sl.shipment_id, False)
        k = group_key(inv.fifo_date, h)
        scanned_by_sid_date[sl.shipment_id][k].append(sl)
        scanned_sum_by_sid_date[sl.shipment_id][k] += Decimal(sl.scanned_quantity or 0)
        refs_with_scans.add(inv.reference.strip().upper())

    # 3. İdeal FIFO penceresini yalnızca gerçekten okutması olan referanslar için hesapla
    ideal: dict[int, set[datetime]] = {}
    if refs_with_scans:
        all_ref_shipments = (
            db.query(Shipment)
            .filter(
                Shipment.status.in_([ShipmentStatus.ACTIVE, ShipmentStatus.COMPLETED]),
                func.upper(func.trim(Shipment.reference)).in_(refs_with_scans),
            )
            .all()
        )
        ideal = ideal_fifo_dates(db, all_ref_shipments)

    # 4. Stoktaki blokajlı etiketleri tüm ilgili referanslar için tek sorguda çek
    ref_keys = {s.reference.strip().upper() for s in shipments}
    blocked_rows = (
        db.query(InventoryLabel)
        .filter(
            InventoryLabel.blocked.is_(True),
            InventoryLabel.hidden.is_(False),
        )
        .order_by(InventoryLabel.fifo_date.asc(), InventoryLabel.id.asc())
        .all()
    )
    blocked_by_ref: dict[str, list[InventoryLabel]] = defaultdict(list)
    blocked_inv_ids: list[int] = []
    for r in blocked_rows:
        rk = r.reference.strip().upper()
        if rk in ref_keys:
            blocked_by_ref[rk].append(r)
            blocked_inv_ids.append(r.id)

    blocked_consumed_ids: set[int] = set()
    if blocked_inv_ids:
        blocked_consumed_ids = {
            row[0]
            for row in db.query(ShipmentLabel.inventory_label_id).filter(
                ShipmentLabel.inventory_label_id.in_(blocked_inv_ids),
                ShipmentLabel.status.in_([ShipmentLabelStatus.SCANNED, ShipmentLabelStatus.PARTIAL]),
            ).all()
        }

    manifests = []
    for s in shipments:
        prog = prog_by_sid[s.id]
        use_hourly = hourly_by_sid[s.id]
        quotas = quotas_by_sid.get(s.id, {})
        inside = ideal.get(s.id, set())
        scanned_by_date = scanned_by_sid_date.get(s.id, {})
        scanned_sums = scanned_sum_by_sid_date.get(s.id, {})

        ref_blocked_rows = blocked_by_ref.get(s.reference.strip().upper(), [])
        blocked_labels = [
            {
                "label": r.label,
                "quantity": float(r.quantity),
                "fifo_time": r.fifo_date.strftime("%d.%m.%Y %H:%M"),
                "blocked": True,
            }
            for r in ref_blocked_rows
            if r.id not in blocked_consumed_ids and Decimal(r.quantity) > 0
        ]
        blocked_by_date: dict = {}
        for r in ref_blocked_rows:
            if Decimal(r.quantity) <= 0:
                continue
            bkey = group_key(r.fifo_date, use_hourly)
            blocked_by_date.setdefault(bkey, []).append(r)

        items = []
        for key in sorted(quotas):
            quota = float(quotas[key])
            done = float(scanned_sums.get(key, Decimal("0")))
            filled = done >= quota - 0.001
            outside = key not in inside and done > 0.001
            label_rows = scanned_by_date.get(key, [])
            b_rows = blocked_by_date.get(key, [])
            items.append({
                "label": format_group(key, use_hourly),
                "reference": s.reference,
                "quantity": quota,
                "scanned_quantity": done,
                "fifo_date": format_group(key, use_hourly),
                "fifo_group_date": format_group(key, use_hourly),
                "status": "scanned" if filled else "pending",
                "is_scanned": filled,
                "outside_fifo": outside,
                "outside_reason": "window" if outside else None,
                "has_blocked": len(b_rows) > 0,
                "blocked_count": len(b_rows),
                "scanned_labels": [
                    {
                        "label": sl.inventory_label.label,
                        "quantity": float(sl.scanned_quantity or 0),
                        "outside_fifo": outside,
                        "blocked": bool(getattr(sl.inventory_label, "blocked", False)),
                    }
                    for sl in label_rows
                    if sl.inventory_label
                ],
            })

        manifests.append({
            "shipment_id": s.id,
            "group_id": s.group_id or s.id,
            "group_name": s.group_name or s.name,
            "reference": s.reference,
            "requested_quantity": prog["requested_quantity"],
            "pool_quantity": prog["pool_quantity"],
            "scanned_quantity": prog["scanned_quantity"],
            "remaining_quantity": prog["remaining_quantity"],
            "progress_percent": prog["progress_percent"],
            "hourly_fifo": use_hourly,
            "status": prog["status"],
            "is_complete": prog["is_complete"],
            "blocked_label_count": prog.get("blocked_label_count", 0),
            "unblocked_label_count": prog.get("unblocked_label_count", 0),
            "blocked_labels": blocked_labels,
            "items": items,
        })

    return manifests


def export_manifest_excel(db: Session, group_id: int | None = None) -> bytes:
    """
    Sevkiyat manifestini ve okutulan tüm etiketleri çok sayfalı Excel (.xlsx) olarak üretir.
    1. Sayfa: Sevkiyat Özeti (Sevkiyat, Referans, İstenen, Okutulan, Kalan, Eksik Stok, Durum)
    2. Sayfa: Okutulan Etiketler (Sevkiyat, Referans, Etiket No, Okutulan Adet, FIFO Tarihi, Durum)
    3. Sayfa: FIFO Tarih Planı (Sevkiyat, Referans, FIFO Tarihi, Beklenen, Okutulan, Durum)
    """
    from io import BytesIO
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    manifests = get_shipment_manifest(db)
    if group_id is not None:
        manifests = [m for m in manifests if (m["group_id"] or m["shipment_id"]) == group_id]

    wb = Workbook()
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")

    # 1. Sayfa: Sevkiyat Özeti
    ws_summary = wb.active
    ws_summary.title = "Sevkiyat Özeti"
    summary_headers = [
        "SEVKİYAT",
        "REFERANS",
        "İSTENEN MİKTAR",
        "OKUTULAN MİKTAR",
        "KALAN MİKTAR",
        "EKSİK STOK",
        "İLERLEME (%)",
        "DURUM",
    ]
    ws_summary.append(summary_headers)
    for cell in ws_summary[1]:
        cell.font = header_font
        cell.fill = header_fill

    for m in manifests:
        prog = get_shipment_progress(db, m["shipment_id"])
        short_q = prog.get("short_quantity", 0)
        status_txt = (
            "Tamamlandı"
            if m["is_complete"] or m["scanned_quantity"] >= m["requested_quantity"]
            else ("Eksik Stok" if short_q > 0 else "Devam Ediyor")
        )
        ws_summary.append([
            m.get("group_name") or f"{m['group_id']}. Sevkiyat",
            m["reference"],
            m["requested_quantity"],
            m["scanned_quantity"],
            m["remaining_quantity"],
            short_q,
            m["progress_percent"],
            status_txt,
        ])

    # 2. Sayfa: Okutulan Etiketler (Packing List / İrsaliye Detayı)
    ws_labels = wb.create_sheet(title="Okutulan Etiketler")
    label_headers = [
        "SEVKİYAT",
        "REFERANS",
        "ETİKET NO",
        "OKUTULAN MİKTAR",
        "FIFO TARİHİ",
        "FIFO DURUMU",
    ]
    ws_labels.append(label_headers)
    for cell in ws_labels[1]:
        cell.font = header_font
        cell.fill = header_fill

    for m in manifests:
        gname = m.get("group_name") or f"{m['group_id']}. Sevkiyat"
        for item in m["items"]:
            for read in item.get("scanned_labels", []):
                ws_labels.append([
                    gname,
                    m["reference"],
                    read["label"],
                    read["quantity"],
                    item["fifo_group_date"],
                    "FIFO Dışı" if read.get("outside_fifo") else "Uygun",
                ])

    # 3. Sayfa: FIFO Tarih Planı
    ws_plan = wb.create_sheet(title="FIFO Tarih Planı")
    plan_headers = [
        "SEVKİYAT",
        "REFERANS",
        "FIFO TARİHİ",
        "BEKLENEN MİKTAR",
        "OKUTULAN MİKTAR",
        "OKUTULAN ETİKETLER",
        "DURUM",
    ]
    ws_plan.append(plan_headers)
    for cell in ws_plan[1]:
        cell.font = header_font
        cell.fill = header_fill

    for m in manifests:
        gname = m.get("group_name") or f"{m['group_id']}. Sevkiyat"
        for item in m["items"]:
            scanned = item.get("scanned_quantity", 0)
            filled = item.get("is_scanned") or scanned >= item["quantity"]
            state_txt = "FIFO Dışı" if item.get("outside_fifo") else ("Doldu" if filled else ("Devam" if scanned > 0 else "Açık"))
            lbl_str = ", ".join(f"{r['label']} ({r['quantity']})" for r in item.get("scanned_labels", []))
            ws_plan.append([
                gname,
                m["reference"],
                item["fifo_group_date"],
                item["quantity"],
                scanned,
                lbl_str or "-",
                state_txt,
            ])

    for ws in (ws_summary, ws_labels, ws_plan):
        for col in ws.columns:
            max_len = max((len(str(c.value or "")) for c in col), default=10)
            ws.column_dimensions[col[0].column_letter].width = min(max(max_len + 3, 12), 45)

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def list_candidate_labels(db: Session, shipment_id: int, fifo_group: str) -> dict:
    """Bir FIFO tarihinde henüz okutulmamış, bu sevkiyata kabul edilecek etiketler."""
    shipment = db.query(Shipment).filter(Shipment.id == shipment_id).first()
    if shipment is None or shipment.status not in (ShipmentStatus.ACTIVE, ShipmentStatus.COMPLETED):
        raise ValueError("Sevkiyat bulunamadı")

    hourly = bool(shipment.hourly_fifo)
    try:
        parsed = datetime.strptime(fifo_group.strip(), "%d.%m.%Y %H:%M" if hourly else "%d.%m.%Y")
    except ValueError as exc:
        raise ValueError("FIFO tarihi okunamadı") from exc

    key = group_key(parsed, hourly)
    quota = quota_map(db, shipment).get(key, Decimal("0"))
    done = scanned_on_date(db, shipment, key)
    remaining = quota - done
    if remaining < 0:
        remaining = Decimal("0")

    end = key + (timedelta(minutes=1) if hourly else timedelta(days=1))
    from sqlalchemy import func
    rows = (
        db.query(InventoryLabel)
        .filter(
            func.upper(func.trim(InventoryLabel.reference)) == shipment.reference.strip().upper(),
            InventoryLabel.hidden.is_(False),
            InventoryLabel.fifo_date >= key,
            InventoryLabel.fifo_date < end,
        )
        .order_by(InventoryLabel.fifo_date.asc(), InventoryLabel.id.asc())
        .all()
    )
    labels = []
    if remaining > 0 and rows:
        consumed = {
            row[0]
            for row in db.query(ShipmentLabel.inventory_label_id).filter(
                ShipmentLabel.inventory_label_id.in_([row.id for row in rows]),
                ShipmentLabel.status.in_([ShipmentLabelStatus.SCANNED, ShipmentLabelStatus.PARTIAL]),
            ).all()
        }
        for inv in rows:
            if inv.id in consumed or Decimal(inv.quantity) <= 0:
                continue
            labels.append({
                "label": inv.label,
                "quantity": float(inv.quantity),
                "fifo_time": inv.fifo_date.strftime("%H:%M"),
                "blocked": bool(getattr(inv, "blocked", False)),
            })

    return {"remaining": float(remaining), "labels": labels}

